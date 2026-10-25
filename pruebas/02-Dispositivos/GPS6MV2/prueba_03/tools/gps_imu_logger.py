"""Registra las salidas del firmware de prueba 03 en un libro Excel."""

from __future__ import annotations

import argparse
from datetime import datetime
from math import isfinite
from pathlib import Path
from time import monotonic

import serial
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


GPS_COLUMNS = [
    "hora_pc", "uptime_esp_us", "tipo", "checksum_valido", "parse_valido",
    "hora_gps_ms", "rmc_valido", "calidad_fix", "desfase_rmc_gga_ms",
    "latitud_deg", "longitud_deg", "velocidad_nudos", "velocidad_kmh",
    "rumbo_deg", "satelites", "hdop", "edad_rmc_ms", "edad_gga_ms",
    "estado_gps", "estado_velocidad", "velocidad_utilizable",
    "nmea_crudo", "linea_serial",
]
IMU_COLUMNS = [
    "hora_pc", "tipo", "recibido_esp_us", "timestamp_reporte_us",
    "x", "y", "z", "w", "estado_sensor", "linea_serial",
]
EVENT_COLUMNS = ["hora_pc", "linea_serial"]
COMBINED_COLUMNS = [
    "hora_pc", "numero_muestra", "tiempo_esp_us", "gps_actualizado",
    "gps_edad_ms", "gps_estado", "gps_calidad_fix", "gps_satelites",
    "gps_hdop", "gps_latitud_deg", "gps_longitud_deg", "gps_velocidad_kmh",
    "gps_rumbo_deg", "imu_roll_deg", "imu_pitch_deg", "imu_yaw_deg",
    "imu_ax_m_s2", "imu_ay_m_s2", "imu_az_m_s2", "estado_orientacion",
    "estado_aceleracion",
]
SAVE_INTERVAL_S = 5.0
COMBINED_INTERVAL_S = 0.5


def _number(value: str, cast=float):
    if value in (None, "", "nan", "NaN"):
        return None
    try:
        result = cast(value)
        if cast is float and not isfinite(result):
            return None
        return result
    except (ValueError, OverflowError):
        return None


def _find_payload(line: str, marker: str) -> str | None:
    index = line.find(marker)
    return line[index:] if index >= 0 else None


def parse_gps(line: str) -> dict | None:
    payload = _find_payload(line, "GPSDIAG|")
    if payload is None:
        return None
    parts = payload.split("|", 26)
    if len(parts) != 27 or parts[1] != "1":
        return None
    (_marker, _version, uptime, sentence_type, checksum, parsed, gps_ms,
     rmc_valid, quality, epoch_skew, latitude, longitude, speed_knots,
     speed_kmh, course, satellites, hdop, rmc_age, gga_age, gps_state,
     speed_state, speed_usable, _delta, _acceleration, _coordinate_speed,
     _coordinate_delta, raw) = parts
    return dict(zip(GPS_COLUMNS, [
        datetime.now().astimezone().isoformat(timespec="milliseconds"),
        _number(uptime, int), sentence_type, _number(checksum, int),
        _number(parsed, int), _number(gps_ms, int), _number(rmc_valid, int),
        _number(quality, int), _number(epoch_skew, int), _number(latitude),
        _number(longitude), _number(speed_knots), _number(speed_kmh),
        _number(course), _number(satellites, int), _number(hdop),
        _number(rmc_age, int), _number(gga_age, int), gps_state, speed_state,
        _number(speed_usable, int), raw, line,
    ]))


def parse_imu(line: str) -> dict | None:
    payload = _find_payload(line, "IMU|")
    if payload is None:
        return None
    parts = payload.split("|")
    if len(parts) not in (9, 10) or parts[1] != "1":
        return None
    # IMU|v|callback_us|SH2_timestamp_us|tipo|x|y|z|[w]|accuracy
    if len(parts) == 9:
        _, _, received, report_time, sensor_type, x, y, z, status = parts
        w = None
    else:
        _, _, received, report_time, sensor_type, x, y, z, w, status = parts
    return dict(zip(IMU_COLUMNS, [
        datetime.now().astimezone().isoformat(timespec="milliseconds"),
        sensor_type, _number(received, int), _number(report_time, int),
        _number(x), _number(y), _number(z), _number(w),
        _number(status, int), line,
    ]))


def _prepare_sheet(book, name: str, columns: list[str]):
    sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
    if sheet.max_row == 0 or not any(cell.value for cell in sheet[1]):
        for index, column in enumerate(columns, start=1):
            sheet.cell(row=1, column=index, value=column)
    elif [cell.value for cell in sheet[1]] != columns:
        raise ValueError(f"La hoja {name} ya existe con encabezados distintos.")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Guarda GPS e IMU del firmware de prueba 03 en Excel."
    )
    parser.add_argument("puerto", help="Puerto serial del ESP32, por ejemplo COM5")
    parser.add_argument("archivo", help="Nombre del archivo .xlsx que se creará")
    args = parser.parse_args()

    output = Path(args.archivo).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        parser.error("El archivo de salida debe terminar en .xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        book = load_workbook(output)
    else:
        book = Workbook()
        book.remove(book.active)
    gps_sheet = _prepare_sheet(book, "GPS", GPS_COLUMNS)
    imu_sheet = _prepare_sheet(book, "IMU", IMU_COLUMNS)
    event_sheet = _prepare_sheet(book, "Eventos", EVENT_COLUMNS)
    combined_sheet = _prepare_sheet(book, "Muestras_500ms", COMBINED_COLUMNS)
    book.save(output)

    last_save = monotonic()
    next_combined = monotonic() + COMBINED_INTERVAL_S
    latest_device_us = None
    latest_rmc = None
    latest_gga = None
    latest_orientation = None
    latest_acceleration = None
    last_combined_rmc_us = None
    print(f"Registro prueba 03: puerto {args.puerto} -> {output}")
    print("Se guardan GPS e IMU crudos y una muestra combinada cada 500 ms.")
    print("Ctrl+C detiene y guarda. La muestra combinada no fusiona ni genera alarmas.")
    try:
        with serial.Serial(args.puerto, 115200, timeout=0.05) as port:
            while True:
                raw = port.readline()
                if raw:
                    line = raw.decode("utf-8", errors="replace").strip()
                    gps_record = parse_gps(line)
                    imu_record = parse_imu(line)
                    if gps_record:
                        gps_sheet.append([gps_record[column] for column in GPS_COLUMNS])
                        latest_device_us = max(latest_device_us or 0, gps_record["uptime_esp_us"] or 0)
                        if gps_record["tipo"] == "GPRMC":
                            latest_rmc = gps_record
                        elif gps_record["tipo"] == "GPGGA":
                            latest_gga = gps_record
                    elif imu_record:
                        imu_sheet.append([imu_record[column] for column in IMU_COLUMNS])
                        latest_device_us = max(latest_device_us or 0, imu_record["recibido_esp_us"] or 0)
                        if imu_record["tipo"] == "ROTATION_EULER":
                            latest_orientation = imu_record
                        elif imu_record["tipo"] == "LINEAR_ACCEL":
                            latest_acceleration = imu_record
                    elif ("GPS_EVENT|" in line or "BNO085" in line
                          or ("IMU" in line and "IMU|" not in line)):
                        event_sheet.append([
                            datetime.now().astimezone().isoformat(timespec="milliseconds"),
                            line,
                        ])
                    else:
                        line = None
                    if line is not None:
                        print(line)

                now = monotonic()
                while latest_device_us is not None and now >= next_combined:
                    gps_age_ms = None
                    if latest_rmc is not None and latest_rmc["uptime_esp_us"] is not None:
                        gps_age_ms = max(0, latest_device_us - latest_rmc["uptime_esp_us"]) / 1000
                    gps_updated = (
                        latest_rmc is not None
                        and latest_rmc["uptime_esp_us"] != last_combined_rmc_us
                    )
                    combined_sheet.append([
                        datetime.now().astimezone().isoformat(timespec="milliseconds"),
                        combined_sheet.max_row - 1,
                        latest_device_us,
                        int(gps_updated),
                        gps_age_ms,
                        latest_rmc["estado_gps"] if latest_rmc else None,
                        latest_gga["calidad_fix"] if latest_gga else None,
                        latest_gga["satelites"] if latest_gga else None,
                        latest_gga["hdop"] if latest_gga else None,
                        latest_rmc["latitud_deg"] if latest_rmc else None,
                        latest_rmc["longitud_deg"] if latest_rmc else None,
                        latest_rmc["velocidad_kmh"] if latest_rmc else None,
                        latest_rmc["rumbo_deg"] if latest_rmc else None,
                        latest_orientation["x"] if latest_orientation else None,
                        latest_orientation["y"] if latest_orientation else None,
                        latest_orientation["z"] if latest_orientation else None,
                        latest_acceleration["x"] if latest_acceleration else None,
                        latest_acceleration["y"] if latest_acceleration else None,
                        latest_acceleration["z"] if latest_acceleration else None,
                        latest_orientation["estado_sensor"] if latest_orientation else None,
                        latest_acceleration["estado_sensor"] if latest_acceleration else None,
                    ])
                    if latest_rmc is not None:
                        last_combined_rmc_us = latest_rmc["uptime_esp_us"]
                    next_combined += COMBINED_INTERVAL_S

                if monotonic() - last_save >= SAVE_INTERVAL_S:
                    book.save(output)
                    last_save = monotonic()
    except KeyboardInterrupt:
        print("\nRegistro detenido por la persona.")
    except serial.SerialException as error:
        print(f"No se pudo abrir {args.puerto}: {error}")
        print("Cierre otros monitores que estén usando el puerto COM e inténtelo de nuevo.")
    finally:
        book.save(output)
        print(f"Archivo guardado: {output}")


if __name__ == "__main__":
    main()
