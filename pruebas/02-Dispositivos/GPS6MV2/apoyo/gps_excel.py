"""Funciones para crear y actualizar el libro de resultados GPS."""

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from gps_parser import COLUMNS


def prepare_workbook(path: Path):
    """Abre un Excel existente o crea uno nuevo con la hoja GPS."""
    if path.exists():
        try:
            book = load_workbook(path)
            sheet = book["GPS"] if "GPS" in book.sheetnames else book.active
            sheet.title = "GPS"
        except (BadZipFile, OSError):
            backup = path.with_name(path.stem + "_corrupto.xlsx")
            try:
                path.replace(backup)
                print(f"El archivo existente no era válido; se conservó como {backup.name}")
                return prepare_workbook(path)
            except PermissionError:
                raise PermissionError(
                    f"{path} está bloqueado. Cierra Excel, su vista previa y OneDrive; "
                    "después vuelve a ejecutar el programa."
                )
        headers = [cell.value for cell in sheet[1]] if sheet.max_row else []
        if not any(headers):
            for index, column in enumerate(COLUMNS, start=1):
                sheet.cell(row=1, column=index, value=column)
            headers = list(COLUMNS)
        for column in COLUMNS:
            if column not in headers:
                sheet.cell(row=1, column=len(headers) + 1, value=column)
                headers.append(column)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(headers)).column_letter}1"
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "GPS"
        sheet.append(COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(COLUMNS)).column_letter}1"
    widths = {
        "hora_pc": 25, "hora_gps": 14, "tipo": 12, "nmea_crudo": 55,
        "estado_gps": 18, "estado_velocidad": 26, "evento": 45,
    }
    for index, column in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = widths.get(column, 16)
    return book, sheet