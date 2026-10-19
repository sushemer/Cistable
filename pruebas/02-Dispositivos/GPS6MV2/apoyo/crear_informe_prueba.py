"""Genera un libro consolidado con datos originales y fórmulas internas."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_ROW = 10
DATA_ROW = HEADER_ROW + 1
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"


def safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name).strip(" '") or "Datos"
    base = cleaned[:31]
    candidate = base
    number = 2
    while candidate.lower() in used:
        suffix = f"_{number}"
        candidate = base[: 31 - len(suffix)] + suffix
        number += 1
    used.add(candidate.lower())
    return candidate


def style_header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def formula_for_gps(headers: list[str], data_row: int, header: str) -> str:
    positions = {str(value): index + 1 for index, value in enumerate(headers)}

    def cell(field: str) -> str | None:
        index = positions.get(field)
        return f"${get_column_letter(index)}{data_row}" if index else None

    kind = cell("tipo")
    if not kind:
        return '=0'
    rmc = f'OR({kind}="GPRMC",{kind}="GNRMC")'
    flags = []
    for field in ("checksum_valido", "parse_valido"):
        reference = cell(field)
        if reference:
            flags.append(f"OR({reference}=1,{reference}=TRUE)")
    fix = cell("rmc_valido") or cell("fix_valido")
    if fix:
        flags.append(f"OR({fix}=1,{fix}=TRUE)")
    valid = f"AND({rmc},{','.join(flags)})" if flags else rmc
    speed = cell("velocidad_kmh")
    usable = cell("velocidad_utilizable")

    if header == "RMC_fix_valido":
        return f"=IF({valid},1,0)"
    if header == "velocidad_RMC_con_fix_kmh":
        return f'=IF(AND({valid},ISNUMBER({speed or ""})),{speed or ""},"")'
    if header == "velocidad_aceptada_firmware_kmh":
        if not usable or not speed:
            return '=IF(FALSE,"","")'
        return f'=IF(AND({valid},OR({usable}=1,{usable}=TRUE),ISNUMBER({speed})),{speed},"")'
    return ""


def make_gps_sheet(wb, used, run_label, source_file, headers, rows):
    ws = wb.create_sheet(safe_sheet_name(f"{run_label}_GPS", used))
    title = f"GPS — {source_file}"
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)

    for col, label in enumerate([
        "Registros GPS", "Registros RMC", "RMC con fix válido",
        "Velocidades RMC válidas", "Promedio (km/h)", "Máximo (km/h)",
        "Velocidades aceptadas", "Duración GPS (s)",
    ], start=1):
        ws.cell(3, col, label)
    style_header(ws, 3, 8)

    columns = list(headers) + [
        "RMC_fix_valido", "velocidad_RMC_con_fix_kmh",
        "velocidad_aceptada_firmware_kmh",
    ]
    header_map = {str(value): index + 1 for index, value in enumerate(headers)}
    last_row = DATA_ROW + len(rows) - 1

    def range_for(field: str) -> str | None:
        column = header_map.get(field)
        return f"${get_column_letter(column)}${DATA_ROW}:${get_column_letter(column)}${last_row}" if column and rows else None

    type_range = range_for("tipo")
    fix_values = f"{get_column_letter(len(headers) + 1)}${DATA_ROW}:{get_column_letter(len(headers) + 1)}${last_row}"
    valid_speed = f"{get_column_letter(len(headers) + 2)}${DATA_ROW}:{get_column_letter(len(headers) + 2)}${last_row}"
    accepted_speed = f"{get_column_letter(len(headers) + 3)}${DATA_ROW}:{get_column_letter(len(headers) + 3)}${last_row}"
    uptime_range = range_for("uptime_esp_us")

    formulas = [
        f"=COUNTA({type_range})" if type_range else "=0",
        f'=COUNTIF({type_range},"*RMC")' if type_range else "=0",
        f'=COUNTIF({fix_values},1)' if rows else "=0",
        f"=COUNT({valid_speed})" if rows else "=0",
        f'=IFERROR(AVERAGE({valid_speed}),"")' if rows else '=IFERROR(AVERAGE(A1:A1),"")',
        f'=IF(COUNT({valid_speed})=0,"",MAX({valid_speed}))' if rows else '=IF(COUNT(A1:A1)=0,"",MAX(A1:A1))',
        f"=COUNT({accepted_speed})" if rows else "=0",
        f"=IFERROR((MAX({uptime_range})-MIN({uptime_range}))/1000000,0)" if uptime_range else "=0",
    ]
    for col, formula in enumerate(formulas, start=1):
        ws.cell(4, col, formula)
        ws.cell(4, col).number_format = "0.000"

    for col, value in enumerate(columns, start=1):
        ws.cell(HEADER_ROW, col, value)
    style_header(ws, HEADER_ROW, len(columns))
    for row_index, row in enumerate(rows, start=DATA_ROW):
        values = list(row[: len(headers)])
        values += [None] * (len(headers) - len(values))
        ws.append([]) if ws.max_row < row_index else None
        for col_index, value in enumerate(values, start=1):
            ws.cell(row_index, col_index, value)
        for offset, derived in enumerate(columns[len(headers):], start=len(headers) + 1):
            ws.cell(row_index, offset, formula_for_gps(headers, row_index, derived))

    ws.freeze_panes = f"A{DATA_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{max(HEADER_ROW,last_row)}"
    ws.column_dimensions["A"].width = 27
    for col in range(2, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    if rows and "uptime_esp_us" in header_map:
        chart = LineChart()
        chart.title = "Velocidad RMC con fix"
        chart.y_axis.title = "km/h"
        chart.x_axis.title = "Tiempo ESP32 (µs)"
        chart.height, chart.width = 8, 16
        data = Reference(ws, min_col=len(headers) + 2, min_row=HEADER_ROW, max_row=last_row)
        xvalues = Reference(ws, min_col=header_map["uptime_esp_us"], min_row=DATA_ROW, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(xvalues)
        ws.add_chart(chart, "J3")
    return ws.title


def make_imu_sheet(wb, used, run_label, source_file, headers, rows):
    ws = wb.create_sheet(safe_sheet_name(f"{run_label}_IMU", used))
    ws.merge_cells("A1:H1")
    ws["A1"] = f"IMU — {source_file}"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    metric_labels = [
        "Muestras IMU", "Aceleración lineal", "Giroscopio", "Cuaterniones",
        "Aceleración media (m/s²)", "Giro medio (rad/s)", "Norma media cuaternión", "Duración IMU (s)",
    ]
    for col, label in enumerate(metric_labels, start=1):
        ws.cell(3, col, label)
    style_header(ws, 3, 8)

    header_map = {str(value): index + 1 for index, value in enumerate(headers)}
    derived = ["tiempo_desde_inicio_s", "magnitud_aceleracion_lineal_m_s2", "magnitud_giroscopio_rad_s", "norma_cuaternion"]
    columns = list(headers) + derived
    last_row = DATA_ROW + len(rows) - 1
    type_col = get_column_letter(header_map.get("tipo", 2))
    report_col = get_column_letter(header_map.get("timestamp_reporte_us", 4))
    x_col = get_column_letter(header_map.get("x", 5))
    y_col = get_column_letter(header_map.get("y", 6))
    z_col = get_column_letter(header_map.get("z", 7))
    w_col = get_column_letter(header_map.get("w", 8))
    type_range = f"${type_col}${DATA_ROW}:${type_col}${last_row}" if rows else None
    accel_mag_col = get_column_letter(len(headers) + 2)
    gyro_mag_col = get_column_letter(len(headers) + 3)
    quat_norm_col = get_column_letter(len(headers) + 4)
    accel_range = f"${accel_mag_col}${DATA_ROW}:${accel_mag_col}${last_row}"
    gyro_range = f"${gyro_mag_col}${DATA_ROW}:${gyro_mag_col}${last_row}"
    quat_range = f"${quat_norm_col}${DATA_ROW}:${quat_norm_col}${last_row}"
    formulas = [
        f"=COUNTA({type_range})" if type_range else "=0",
        f'=COUNTIF({type_range},"LINEAR_ACCEL")' if type_range else "=0",
        f'=COUNTIF({type_range},"GYRO")' if type_range else "=0",
        f'=COUNTIF({type_range},"ROTATION_QUAT")' if type_range else "=0",
        f'=IFERROR(AVERAGE({accel_range}),"")' if rows else '=IFERROR(AVERAGE(A1:A1),"")',
        f'=IFERROR(AVERAGE({gyro_range}),"")' if rows else '=IFERROR(AVERAGE(A1:A1),"")',
        f'=IFERROR(AVERAGE({quat_range}),"")' if rows else '=IFERROR(AVERAGE(A1:A1),"")',
        f"=IFERROR((MAX(${report_col}${DATA_ROW}:${report_col}${last_row})-MIN(${report_col}${DATA_ROW}:${report_col}${last_row}))/1000000,0)" if rows else "=0",
    ]
    for col, formula in enumerate(formulas, start=1):
        ws.cell(4, col, formula)
        ws.cell(4, col).number_format = "0.000"

    for col, value in enumerate(columns, start=1):
        ws.cell(HEADER_ROW, col, value)
    style_header(ws, HEADER_ROW, len(columns))
    for row_index, row in enumerate(rows, start=DATA_ROW):
        values = list(row[: len(headers)])
        values += [None] * (len(headers) - len(values))
        for col_index, value in enumerate(values, start=1):
            ws.cell(row_index, col_index, value)
        ws.cell(row_index, len(headers) + 1,
                f"=IF(ISNUMBER(${report_col}{row_index}),(${report_col}{row_index}-${report_col}${DATA_ROW})/1000000,\"\")")
        ws.cell(row_index, len(headers) + 2,
                f'=IF(${type_col}{row_index}="LINEAR_ACCEL",SQRT(${x_col}{row_index}^2+${y_col}{row_index}^2+${z_col}{row_index}^2),"")')
        ws.cell(row_index, len(headers) + 3,
                f'=IF(${type_col}{row_index}="GYRO",SQRT(${x_col}{row_index}^2+${y_col}{row_index}^2+${z_col}{row_index}^2),"")')
        ws.cell(row_index, len(headers) + 4,
                f'=IF(${type_col}{row_index}="ROTATION_QUAT",SQRT(${x_col}{row_index}^2+${y_col}{row_index}^2+${z_col}{row_index}^2+${w_col}{row_index}^2),"")')

    ws.freeze_panes = f"A{DATA_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{max(HEADER_ROW,last_row)}"
    for col in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    if rows:
        for metric_col, title, unit, anchor in [
            (len(headers) + 2, "Aceleración lineal", "m/s²", "J3"),
            (len(headers) + 3, "Velocidad angular giroscopio", "rad/s", "J19"),
        ]:
            chart = LineChart()
            chart.title, chart.y_axis.title, chart.x_axis.title = title, unit, "Tiempo desde inicio (s)"
            chart.height, chart.width = 8, 16
            data = Reference(ws, min_col=metric_col, min_row=HEADER_ROW, max_row=last_row)
            xvalues = Reference(ws, min_col=len(headers) + 1, min_row=DATA_ROW, max_row=last_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(xvalues)
            ws.add_chart(chart, anchor)
    return ws.title


def add_event_sheet(wb, used, run_label, source_file, headers, rows):
    ws = wb.create_sheet(safe_sheet_name(f"{run_label}_Eventos", used))
    ws.merge_cells("A1:C1")
    ws["A1"] = f"Eventos — {source_file}"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A3"] = "Registros de evento"
    ws["B3"] = f"=COUNTA(A{DATA_ROW}:A{DATA_ROW + len(rows) - 1})" if rows else "=0"
    for col, value in enumerate(headers, start=1):
        ws.cell(HEADER_ROW, col, value)
    style_header(ws, HEADER_ROW, max(1, len(headers)))
    for row_index, row in enumerate(rows, start=DATA_ROW):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row_index, col_index, value)
    ws.freeze_panes = f"A{DATA_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(max(1,len(headers)))}{max(HEADER_ROW,DATA_ROW+len(rows)-1)}"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    return ws.title


def consolidate(test_dir: Path, excludes: set[str], force: bool, exclusion_note: str) -> Path:
    test_dir = test_dir.resolve()
    source_dir = test_dir / "pruebas"
    output_dir = test_dir / "analisis"
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{test_dir.name}_resultados.xlsx"
    if output.exists() and not force:
        raise FileExistsError(f"Ya existe {output}. Usa --force para reemplazarlo.")

    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    summary.merge_cells("A1:Q1")
    summary["A1"] = f"Resultados consolidados — {test_dir.name}"
    summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=BLUE)
    summary.merge_cells("A2:Q2")
    summary["A2"] = "Las métricas y totales se calculan mediante fórmulas internas con rangos explícitos; los datos originales se conservan en las hojas por prueba."
    summary["A2"].alignment = Alignment(wrap_text=True)
    headers = [
        "N.º", "Archivo original", "GPS: filas", "GPS: RMC con fix", "GPS: velocidades RMC",
        "GPS: media km/h", "GPS: máximo km/h", "GPS: aceptadas firmware",
        "IMU: muestras", "IMU: aceleración muestras", "IMU: aceleración media m/s²",
        "IMU: giroscopio muestras", "IMU: giro medio rad/s", "IMU: cuaterniones",
        "GPS: duración s", "IMU: duración s", "Eventos",
    ]
    for col, name in enumerate(headers, start=1):
        summary.cell(6, col, name)
    style_header(summary, 6, len(headers))

    sources = wb.create_sheet("Fuentes")
    sources.append(["Archivo de entrada", "Estado", "Hojas importadas", "Nota"])
    style_header(sources, 1, 4)
    used_names = {"resumen", "fuentes"}
    report_rows = []
    files = sorted(source_dir.glob("*.xlsx")) if source_dir.exists() else []
    for source in files:
        if source.name in excludes:
            sources.append([source.name, "No incluido", "", exclusion_note])
            continue
        try:
            src_wb = load_workbook(source, read_only=True, data_only=True)
        except Exception as error:
            sources.append([source.name, "Error al leer", "", str(error)[:300]])
            continue

        label = source.stem.removeprefix(f"{test_dir.name}_")
        imported_tabs = []
        tab_refs = {"GPS": None, "IMU": None, "Eventos": None}
        for src_ws in src_wb.worksheets:
            iterator = src_ws.iter_rows(values_only=True)
            headers_row = next(iterator, None)
            if not headers_row or not any(value is not None for value in headers_row):
                continue
            headers_row = list(headers_row)
            rows = [tuple(row) for row in iterator if any(value is not None for value in row)]
            kind = src_ws.title.strip().upper()
            if kind == "GPS":
                tab_refs["GPS"] = make_gps_sheet(wb, used_names, label, source.name, headers_row, rows)
                imported_tabs.append(tab_refs["GPS"])
            elif kind == "IMU":
                tab_refs["IMU"] = make_imu_sheet(wb, used_names, label, source.name, headers_row, rows)
                imported_tabs.append(tab_refs["IMU"])
            elif kind in ("EVENTOS", "EVENTS"):
                tab_refs["Eventos"] = add_event_sheet(wb, used_names, label, source.name, headers_row, rows)
                imported_tabs.append(tab_refs["Eventos"])
        src_wb.close()
        sources.append([source.name, "Importado" if imported_tabs else "Sin datos", ", ".join(imported_tabs), "Se copiaron los valores de origen; el archivo original permanece intacto."])
        report_rows.append((source.name, tab_refs))

    first_summary_row = 7
    for offset, (filename, refs) in enumerate(report_rows):
        row = first_summary_row + offset
        summary.cell(row, 1, offset + 1)
        summary.cell(row, 2, filename)
        gps_ref, imu_ref, event_ref = refs["GPS"], refs["IMU"], refs["Eventos"]
        # Each summary metric points to the formula cell on that run's detail sheet.
        mapping = {
            3: (gps_ref, "A4"), 4: (gps_ref, "C4"), 5: (gps_ref, "D4"),
            6: (gps_ref, "E4"), 7: (gps_ref, "F4"), 8: (gps_ref, "G4"),
            9: (imu_ref, "A4"), 10: (imu_ref, "B4"), 11: (imu_ref, "E4"),
            12: (imu_ref, "C4"), 13: (imu_ref, "F4"), 14: (imu_ref, "D4"),
            15: (gps_ref, "H4"), 16: (imu_ref, "H4"),
            17: (event_ref, "B3"),
        }
        for col, (sheet_name, cell_ref) in mapping.items():
            summary.cell(row, col, f"=IFERROR('{sheet_name}'!{cell_ref},0)" if sheet_name else "=0")

    total_row = first_summary_row + len(report_rows)
    summary.cell(total_row, 2, "TOTAL / CONSOLIDADO")
    last_data_row = total_row - 1
    if report_rows:
        for col in (3, 4, 5, 8, 9, 10, 12, 14, 15, 16, 17):
            letter = get_column_letter(col)
            summary.cell(total_row, col, f"=SUM({letter}{first_summary_row}:{letter}{last_data_row})")
        summary.cell(total_row, 6, f'=IFERROR(SUMPRODUCT(E{first_summary_row}:E{last_data_row},F{first_summary_row}:F{last_data_row})/SUM(E{first_summary_row}:E{last_data_row}),"")')
        summary.cell(total_row, 7, f'=IF(COUNT(G{first_summary_row}:G{last_data_row})=0,"",MAX(G{first_summary_row}:G{last_data_row}))')
        summary.cell(total_row, 11, f'=IFERROR(SUMPRODUCT(J{first_summary_row}:J{last_data_row},K{first_summary_row}:K{last_data_row})/SUM(J{first_summary_row}:J{last_data_row}),"")')
        summary.cell(total_row, 13, f'=IFERROR(SUMPRODUCT(L{first_summary_row}:L{last_data_row},M{first_summary_row}:M{last_data_row})/SUM(L{first_summary_row}:L{last_data_row}),"")')
    for cell in summary[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    summary.freeze_panes = "A7"
    summary.auto_filter.ref = f"A6:Q{max(6,total_row-1)}"
    for col in range(1, 18):
        summary.column_dimensions[get_column_letter(col)].width = 22
    summary.column_dimensions["B"].width = 32
    summary.row_dimensions[2].height = 34

    sources.freeze_panes = "A2"
    sources.column_dimensions["A"].width = 38
    sources.column_dimensions["B"].width = 20
    sources.column_dimensions["C"].width = 45
    sources.column_dimensions["D"].width = 75
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except AttributeError:
        pass
    wb.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("prueba", choices=("prueba_02", "prueba_03"), help="Carpeta de resultados")
    parser.add_argument("--excluir", action="append", default=[], metavar="ARCHIVO.xlsx", help="Omitir fuente no disponible (se anota en Fuentes)")
    parser.add_argument("--motivo-exclusion", default="Omitido explícitamente; el original no se modificó.", help="Motivo que aparecerá en la hoja Fuentes")
    parser.add_argument("--force", action="store_true", help="Reemplazar el informe existente")
    args = parser.parse_args()
    gps_root = Path(__file__).resolve().parents[1]
    test_dir = gps_root / args.prueba
    try:
        output = consolidate(test_dir, set(args.excluir), args.force, args.motivo_exclusion)
    except FileExistsError as error:
        parser.error(str(error))
    print(f"Informe creado: {output}")


if __name__ == "__main__":
    main()
