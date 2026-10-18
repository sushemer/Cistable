"""Registra datos NMEA del ESP32 en un libro Excel.

Uso: py gps_logger_excel.py --port COM7 --output registro.xlsx
"""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import serial
from serial.serialutil import SerialException
from openpyxl import Workbook, load_workbook
from zipfile import BadZipFile
from openpyxl.styles import Font, PatternFill

# Encabezados que se crearán en la hoja GPS del archivo Excel.
COLUMNS = ["hora_pc", "hora_gps", "tipo", "fix_valido", "latitud_deg", "longitud_deg", "velocidad_kmh", 
           "rumbo_deg", "satelites","hdop", "calidad_fix", "nmea_crudo"]

SERIAL_BAUDRATE = 115200
SAVE_EVERY_ROWS = 10

def nmea_to_degrees(value, hemisphere):
    """Convierte coordenadas NMEA de grados/minutos a grados decimales."""
    if not value or not hemisphere:
        return None
    try:
        raw = float(value)
        result = int(raw / 100) + (raw % 100) / 60
        return -result if hemisphere in ("S", "W") else result
    except ValueError:
        return None

def parse_nmea(line):
    """Interpreta una línea NMEA y la convierte en un registro para Excel."""
    # Se separan los campos por coma para poder identificar cada dato.
    fields = line.split(",")
    # Se inicia el registro con todas las columnas vacías.
    record = {column: None for column in COLUMNS}
    # Esta hora la genera la computadora, no viene del GPS.
    record["hora_pc"] = datetime.now().astimezone().isoformat(timespec="milliseconds")
    # Se conserva la línea recibida, después de quitar espacios y salto de línea.
    record["nmea_crudo"] = line
    if not fields or not fields[0].startswith("$"):
        return record
    # Se guarda el tipo sin el signo inicial $, por ejemplo GNRMC o GNGGA.
    record["tipo"] = fields[0][1:]
    # El checksum queda fuera de los campos interpretados; la línea original no cambia.
    fields[-1] = fields[-1].split("*")[0]
    if fields[0] in ("$GPRMC", "$GNRMC") and len(fields) >= 9:
        # RMC: hora, estado, posición, velocidad y rumbo.
        record.update(hora_gps=fields[1], fix_valido=fields[2] == "A",
                      latitud_deg=nmea_to_degrees(fields[3], fields[4]),
                      longitud_deg=nmea_to_degrees(fields[5], fields[6]))
        try:
            # RMC entrega velocidad en nudos; 1 nudo = 1.852 km/h.
            record["velocidad_kmh"] = float(fields[7] or 0) * 1.852
            record["rumbo_deg"] = float(fields[8] or 0)
        except ValueError:
            pass
    elif fields[0] in ("$GPGGA", "$GNGGA") and len(fields) >= 9:
        # GGA: hora, calidad del fix, satélites y HDOP.
        record["hora_gps"] = fields[1]
        try:
            record["calidad_fix"] = int(fields[6] or 0)
            record["satelites"] = int(fields[7] or 0)
            record["hdop"] = float(fields[8]) if fields[8] else None
        except ValueError:
            pass
    return record

def prepare_workbook(path):
    """Abre un Excel existente o crea uno nuevo con la hoja GPS."""
    if path.exists():
        try:
            # Si el archivo es válido, se continúa agregando registros a su hoja GPS.
            book = load_workbook(path)
            sheet = book["GPS"] if "GPS" in book.sheetnames else book.active
            sheet.title = "GPS"
        except (BadZipFile, OSError):
            # Un archivo dañado se conserva con otro nombre antes de crear uno nuevo.
            backup = path.with_name(path.stem + "_corrupto.xlsx")
            try:
                path.replace(backup)
                print(f"El archivo existente no era válido; se conservó como {backup.name}")
                book, sheet = prepare_workbook(path)
            except PermissionError:
                raise PermissionError(
                    f"{path} está bloqueado. Cierra Excel, su vista previa y OneDrive; "
                    "después vuelve a ejecutar el programa."
                )
    else:
        # Libro nuevo: crea la hoja, encabezados, filtro, congelación y anchos.
        book = Workbook()
        sheet = book.active
        sheet.title = "GPS"
        sheet.append(COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:L1"
        for column, width in zip("ABCDEFGHIJKL", [24, 12, 10, 12, 14, 15, 16, 12, 11, 10, 13, 42]):
            sheet.column_dimensions[column].width = width
    return book, sheet

def main():
    """Recibe líneas del puerto UART y las agrega al Excel indicado.

    Uso: py gps_logger_excel.py --port COM7 --output registro.xlsx
    """
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", required=True, help="COM del ESP32, por ejemplo COM7")
    parser.add_argument("--output", required=True, help="Nombre del archivo Excel de salida")
    args = parser.parse_args()
    # Convierte la ruta de salida en una ruta absoluta.
    output = Path(args.output).resolve()
    book, sheet = prepare_workbook(output)
    book.save(output)
    pending = 0
    print(f"Registrando {args.port} en {output}. Ctrl+C para terminar.")
    try:
        # Abre el puerto COM. El timeout evita quedar bloqueado indefinidamente.
        with serial.Serial(args.port, SERIAL_BAUDRATE, timeout=1) as port:
            while True:
                # Lee bytes directamente del UART hasta encontrar el salto de línea.
                raw = port.readline()
                if not raw:
                    continue
                # Convierte bytes a texto ASCII y elimina espacios/saltos externos.
                line = raw.decode("ascii", errors="replace").strip()
                # Solo se procesan sentencias NMEA que comienzan con '$'.
                if not line.startswith("$"):
                    continue
                # Interpreta la sentencia y la convierte en columnas.
                record = parse_nmea(line)
                # Agrega una fila con los datos interpretados y la línea cruda.
                sheet.append([record[column] for column in COLUMNS])
                pending += 1
                # Muestra en consola la sentencia NMEA ya decodificada y sin el salto final.
                print(line)
                if pending >= SAVE_EVERY_ROWS:
                    # Guarda cada cierto número de filas para reducir la pérdida de datos.
                    book.save(output)
                    pending = 0
    except SerialException as error:
        print(f"No se pudo abrir {args.port}: {error}")
        print("Cierre el monitor, Arduino Serial Monitor, PuTTY o cualquier otra aplicación que use ese COM y vuelve a intentar.")
    except KeyboardInterrupt:
        # Guarda las últimas filas cuando el usuario detiene con Ctrl+C.
        book.save(output)
        print(f"Registro guardado en {output}")

if __name__ == "__main__":
    main()
