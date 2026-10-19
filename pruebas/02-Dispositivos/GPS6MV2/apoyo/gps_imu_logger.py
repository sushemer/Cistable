"""Registra las salidas del firmware de prueba 03 en un libro Excel."""

from __future__ import annotations

import argparse
from datetime import datetime
from math import isfinite
from pathlib import Path
from time import monotonic

import serial
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


GPS_COLUMNS = [
    "hora_pc", "uptime_esp_us", "tipo", "checksum_valido", "parse_valido",
    "hora_gps_ms", "rmc_valido", "calidad_fix", "desfase_rmc_gga_ms",
    "latitud_deg", "longitud_deg", "velocidad_nudos", "velocidad_kmh",
    "rumbo_deg", "satelites", "hdop", "edad_rmc_ms", "edad_gga_ms",
    "estado_gps", "estado_velocidad", "velocidad_utilizable",
    "nmea_crudo", "linea_serial",
]
IMU_COLUMNS = [
    "hora_pc", "tipo", "recibido_esp_us", "timestamp_reporte_us",
    "x", "y", "z", "w", "estado_sensor", "linea_serial",
]
EVENT_COLUMNS = ["hora_pc", "linea_serial"]
SAVE_INTERVAL_S = 5.0


def _number(value: str, cast=float):
    if value in (None, "", "nan", "NaN"):
        return None
    try:
        result = cast(value)
        if cast is float and not isfinite(result):
            return None
        return result
    except (ValueError, OverflowError):
        return None


def _find_payload(line: str, marker: str) -> str | None:
    index = line.find(marker)
    return line[index:] if index >= 0 else None


def parse_gps(line: str) -> dict | None:
    payload = _find_payload(line, "GPSDIAG|")
    if payload is None:
        return None
    parts = payload.split("|", 26)
    if len(parts) != 27 or parts[1] != "1":
        return None
    (_marker, _version, uptime, sentence_type, checksum, parsed, gps_ms,
     rmc_valid, quality, epoch_skew, latitude, longitude, speed_knots,
     speed_kmh, course, satellites, hdop, rmc_age, gga_age, gps_state,
     speed_state, speed_usable, _delta, _acceleration, _coordinate_speed,
     _coordinate_delta, raw) = parts
    return dict(zip(GPS_COLUMNS, [
        datetime.now().astimezone().isoformat(timespec="milliseconds"),
        _number(uptime, int), sentence_type, _number(checksum, int),
        _number(parsed, int), _number(gps_ms, int), _number(rmc_valid, int),
        _number(quality, int), _number(epoch_skew, int), _number(latitude),
        _number(longitude), _number(speed_knots), _number(speed_kmh),
        _number(course), _number(satellites, int), _number(hdop),
        _number(rmc_age, int), _number(gga_age, int), gps_state, speed_state,
        _number(speed_usable, int), raw, line,
    ]))


def parse_imu(line: str) -> dict | None:
    payload = _find_payload(line, "IMU|")
    if payload is None:
        return None
    parts = payload.split("|")
    if len(parts) not in (9, 10) or parts[1] != "1":
        return None
    # IMU|v|callback_us|SH2_timestamp_us|tipo|x|y|z|[w]|accuracy
    if len(parts) == 9:
        _, _, received, report_time, sensor_type, x, y, z, status = parts
        w = None
    else:
        _, _, received, report_time, sensor_type, x, y, z, w, status = parts
    return dict(zip(IMU_COLUMNS, [
        datetime.now().astimezone().isoformat(timespec="milliseconds"),
        sensor_type, _number(received, int), _number(report_time, int),
        _number(x), _number(y), _number(z), _number(w),
        _number(status, int), line,
    ]))


def _prepare_sheet(book, name: str, columns: list[str]):
    sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
    if sheet.max_row == 0 or not any(cell.value for cell in sheet[1]):
        for index, column in enumerate(columns, start=1):
            sheet.cell(row=1, column=index, value=column)
    elif [cell.value for cell in sheet[1]] != columns:
        raise ValueError(f"La hoja {name} ya existe con encabezados distintos.")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Guarda GPS e IMU del firmware de prueba 03 en Excel."
    )
    parser.add_argument("puerto", help="Puerto serial del ESP32, por ejemplo COM5")
    parser.add_argument("archivo", help="Nombre del archivo .xlsx que se creará")
    args = parser.parse_args()

    output = Path(args.archivo).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        parser.error("El archivo de salida debe terminar en .xlsx")
    if output.exists():
        book = load_workbook(output)
    else:
        book = Workbook()
        book.remove(book.active)
    gps_sheet = _prepare_sheet(book, "GPS", GPS_COLUMNS)
    imu_sheet = _prepare_sheet(book, "IMU", IMU_COLUMNS)
    event_sheet = _prepare_sheet(book, "Eventos", EVENT_COLUMNS)
    book.save(output)

    last_save = monotonic()
    print(f"Registro prueba 03: puerto {args.puerto} -> {output}")
    print("GPS e IMU se guardan en hojas separadas. Ctrl+C detiene y guarda.")
    try:
        with serial.Serial(args.puerto, 115200, timeout=1) as port:
            while True:
                raw = port.readline()
                if not raw:
                    continue
                line = raw.decode("utf-8", errors="replace").strip()
                gps_record = parse_gps(line)
                imu_record = parse_imu(line)
                if gps_record:
                    gps_sheet.append([gps_record[column] for column in GPS_COLUMNS])
                elif imu_record:
                    imu_sheet.append([imu_record[column] for column in IMU_COLUMNS])
                elif "GPS_EVENT|" in line or "BNO085" in line or "IMU" in line and "IMU|" not in line:
                    event_sheet.append([
                        datetime.now().astimezone().isoformat(timespec="milliseconds"),
                        line,
                    ])
                else:
                    continue

                print(line)
                if monotonic() - last_save >= SAVE_INTERVAL_S:
                    book.save(output)
                    last_save = monotonic()
    except KeyboardInterrupt:
        print("\nRegistro detenido por la persona.")
    except serial.SerialException as error:
        print(f"No se pudo abrir {args.puerto}: {error}")
        print("Cierre otros monitores que estén usando el puerto COM e inténtelo de nuevo.")
    finally:
        book.save(output)
        print(f"Archivo guardado: {output}")


if __name__ == "__main__":
    main()