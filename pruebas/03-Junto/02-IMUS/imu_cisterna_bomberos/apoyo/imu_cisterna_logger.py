"""Registra el puerto serie, conserva cada línea y separa los campos reconocidos."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from time import monotonic

import serial
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SAVE_SECONDS = 5.0
ANSI = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]")
LOG = re.compile(r"\b([DIWE])\s+\((\d+)\)\s+([^:]+):\s?(.*)")
MASTER_FIELDS = [
    "master_uptime_us", "gps_state", "latitude_deg", "longitude_deg", "gps_speed_kmh",
    "gps_satellites", "gps_hdop", "gps_rmc_age_ms", "front_imu_online", "front_roll_deg",
    "front_pitch_deg", "front_ax_m_s2", "front_ay_m_s2", "front_az_m_s2", "rear_online",
    "rear_sequence", "valid_mask", "imu1_roll_deg", "imu1_pitch_deg", "imu1_ay_m_s2",
    "imu2_roll_deg", "imu2_pitch_deg", "imu2_ay_m_s2", "water_level",
]
DIRECT_IMU_FIELDS = [
    "imu_id", "imu_ubicacion", "imu_sample_time_us", "imu_online", "imu_roll_deg", "imu_pitch_deg", "imu_yaw_deg",
    "imu_ax_m_s2", "imu_ay_m_s2", "imu_az_m_s2", "imu_gx_rad_s", "imu_gy_rad_s", "imu_gz_rad_s",
]
GPS_FIELDS = ["gps_uptime_us", "gps_state", "latitude_deg", "longitude_deg", "gps_speed_kmh", "gps_satellites", "gps_hdop", "gps_speed_valid"]
RSSI_FIELDS = ["rssi_dbm", "rssi_valid", "espnow_sequence"]
RAW_COLUMNS = ["hora_pc_local", "nivel", "etiqueta", "tipo_registro", *DIRECT_IMU_FIELDS,
               *GPS_FIELDS, *RSSI_FIELDS, *MASTER_FIELDS, "nivel_agua", "agua_valida",
               "campos_extra", "campos_separados", "mensaje_crudo", "linea_serial"]


def now_local() -> str:
    return datetime.now().astimezone().isoformat(timespec="milliseconds")


def value(text: str):
    if text.lower() in {"nan", "inf", "+inf", "-inf", "infinity", "+infinity"}:
        return None
    if re.fullmatch(r"[+-]?\d+", text):
        return int(text)
    if re.fullmatch(r"[+-]?(?:\d+\.\d*|\.\d+|\d+)(?:[eE][+-]?\d+)?", text):
        return float(text)
    return text


def make_sheet(book: Workbook):
    ws = book.create_sheet("Datos_Crudos")
    ws.append(RAW_COLUMNS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(RAW_COLUMNS))}1"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.sheet_view.showGridLines = False
    return ws


def fit(ws):
    for column in ws.columns:
        width = min(42, max(12, max(len(str(cell.value or "")) for cell in column[:200]) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    ws.column_dimensions[get_column_letter(RAW_COLUMNS.index("mensaje_crudo") + 1)].width = 55
    ws.column_dimensions[get_column_letter(RAW_COLUMNS.index("linea_serial") + 1)].width = 55


def create_book():
    book = Workbook()
    book.remove(book.active)
    ws = make_sheet(book)
    fit(ws)
    return book, ws


def parse_masterdata(message: str):
    if "MASTERDATA|" not in message:
        return None
    parts = message[message.index("MASTERDATA|"):].split("|")
    if len(parts) < 3 or parts[1] != "2" or len(parts[2:]) != len(MASTER_FIELDS):
        return None
    return dict(zip(MASTER_FIELDS, [value(item) for item in parts[2:]]))


def parse_direct_imu(message: str):
    if "IMU_DATA|" not in message:
        return None
    parts = message[message.index("IMU_DATA|"):].split("|")
    if len(parts) != 13:
        return None
    fields = dict(zip(DIRECT_IMU_FIELDS[:1] + DIRECT_IMU_FIELDS[2:], [value(item) for item in parts[1:]]))
    fields["imu_ubicacion"] = {1: "FRONTAL", 2: "TRASERA"}.get(fields.get("imu_id"), "PENDIENTE")
    return fields


def parse_structured(message: str):
    """Parsea registros conocidos y deja intactos los desconocidos en mensaje_crudo."""
    if "GPS_DATA|" in message:
        parts = message[message.index("GPS_DATA|"):].split("|")
        if len(parts) == 9:
            return "GPS_DATA", dict(zip(GPS_FIELDS, [value(item) for item in parts[1:]]))
    if "RSSI_DATA|" in message:
        parts = message[message.index("RSSI_DATA|"):].split("|")
        if len(parts) == 4:
            return "RSSI_DATA", dict(zip(RSSI_FIELDS, [value(item) for item in parts[1:]]))
    if "WATER_DATA|" in message:
        parts = message[message.index("WATER_DATA|"):].split("|")
        if len(parts) == 4:
            return "WATER_DATA", {"nivel_agua": value(parts[1]), "agua_valida": value(parts[2]), "banderas": parts[3]}
    if "GPSDIAG|" in message or "GPSFUSION|" in message:
        kind = "GPSDIAG" if "GPSDIAG|" in message else "GPSFUSION"
        return kind, {"registro_gps_crudo": message[message.index(kind + "|"): ]}
    direct = parse_direct_imu(message)
    if direct:
        return "IMU_DATA", direct
    master = parse_masterdata(message)
    if master:
        return "MASTERDATA", master
    # Formato flexible: conserva los campos por posición aun si el tipo de
    # mensaje todavía no está documentado (por ejemplo, telemetría de agua).
    payload = message.strip()
    delimiter = "|" if "|" in payload else "," if "," in payload else None
    if delimiter:
        tokens = [item.strip() for item in payload.split(delimiter)]
        return "OTRA_LINEA_SEPARADA", {"campos_separados": tokens, "separador": delimiter}
    return "OTRA_LINEA", {"campos_separados": [payload] if payload else []}


def append_row(ws, level: str, tag: str, message: str, line: str):
    kind, fields = parse_structured(message)
    direct_values = [fields.get(field) for field in DIRECT_IMU_FIELDS]
    gps_values = [fields.get(field) for field in GPS_FIELDS]
    rssi_values = [fields.get(field) for field in RSSI_FIELDS]
    master_values = [fields.get(field) for field in MASTER_FIELDS]
    separated = fields.get("campos_separados")
    ws.append([now_local(), level, tag, kind, *direct_values, *gps_values, *rssi_values,
               *master_values, fields.get("nivel_agua"), fields.get("agua_valida"),
               json.dumps(fields, ensure_ascii=False, default=str) if fields else None,
               json.dumps(separated, ensure_ascii=False, default=str) if separated is not None else None,
               message, line])


def main():
    parser = argparse.ArgumentParser(description="Guarda cada línea serial en crudo y separa los formatos reconocidos; conserva los demás.")
    parser.add_argument("puerto", help="Puerto COM del ESP32, por ejemplo COM5")
    parser.add_argument("archivo", help="Archivo Excel nuevo, por ejemplo registros\\registro_atras.xlsx")
    args = parser.parse_args()
    output = Path(args.archivo).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        parser.error("El archivo de salida debe tener la extensión .xlsx")
    if output.exists():
        parser.error(f"No se reemplazará un archivo existente: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    book, ws = create_book()
    book.save(output)
    print(f"Registrando {args.puerto} en {output}")
    print("Cada fila conserva la línea serial completa. Se separan IMU y GPS conocidos; otros mensajes delimitados se guardan por campos y en crudo.")
    last_save = monotonic()
    try:
        with serial.Serial(args.puerto, 115200, timeout=1) as port:
            while True:
                raw_bytes = port.readline()
                if not raw_bytes:
                    if monotonic() - last_save >= SAVE_SECONDS:
                        book.save(output)
                        last_save = monotonic()
                    continue
                line = ANSI.sub("", raw_bytes.decode("utf-8", errors="replace").rstrip("\r\n"))
                match = LOG.search(line)
                level = match.group(1) if match else ""
                tag = match.group(3).strip() if match else ""
                message = match.group(4) if match else line
                append_row(ws, level, tag, message, line)
                print(line)
                if monotonic() - last_save >= SAVE_SECONDS:
                    book.save(output)
                    last_save = monotonic()
    except KeyboardInterrupt:
        print("\nRegistro detenido por el usuario.")
    except serial.SerialException as error:
        print(f"No se pudo abrir el puerto {args.puerto}: {error}")
    finally:
        fit(ws)
        book.save(output)
        print(f"Archivo Excel guardado: {output}")


if __name__ == "__main__":
    main()
