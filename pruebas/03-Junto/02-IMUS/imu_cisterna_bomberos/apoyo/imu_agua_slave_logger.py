"""Registra la salida actual del slave y separa IMU/agua cuando aparecen."""

from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path
from time import monotonic

import serial
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SAVE_SECONDS = 5.0
LOG = re.compile(r"\b([DIWE])\s+\((\d+)\)\s+([^:]+):\s?(.*)")
SENSORS = re.compile(
    r"Enviando sensores:\s*IMU\[roll=([^,]+),\s*pitch=([^,]+),\s*ay=([^\]]+)\]\s*"
    r"nivel=([^ ]+)\s+estado=0x([0-9A-Fa-f]+)"
)
COLUMNS = ["hora_pc_local", "nivel_log", "etiqueta", "tipo_registro", "imu_roll", "imu_pitch",
           "imu_ay", "nivel_agua", "banderas_hex", "campos_pendientes", "mensaje_crudo", "linea_serial"]


def number(text: str):
    try:
        return float(text) if text.lower() not in {"nan", "inf", "-inf"} else None
    except ValueError:
        return text


def main():
    parser = argparse.ArgumentParser(description="Registra la salida actual del slave y separa IMU y agua cuando aparecen.")
    parser.add_argument("puerto", help="Puerto COM del ESP32 slave, por ejemplo COM7")
    parser.add_argument("archivo", help="Archivo Excel nuevo, por ejemplo registros\\cisterna.xlsx")
    args = parser.parse_args()
    output = Path(args.archivo).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        parser.error("El archivo de salida debe tener extensión .xlsx")
    if output.exists():
        parser.error(f"No se reemplazará un archivo existente: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Registro_slave"
    sheet.append(COLUMNS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:L1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col, width in {"A": 27, "D": 20, "J": 42, "K": 55, "L": 55}.items():
        sheet.column_dimensions[col].width = width
    book.save(output)
    print(f"Registrando {args.puerto} en {output}")

    last_save = monotonic()
    try:
        with serial.Serial(args.puerto, 115200, timeout=1) as port:
            while True:
                raw = port.readline()
                if not raw:
                    if monotonic() - last_save >= SAVE_SECONDS:
                        book.save(output)
                        last_save = monotonic()
                    continue
                line = raw.decode("utf-8", errors="replace").rstrip("\r\n")
                match = LOG.search(line)
                level = match.group(1) if match else ""
                tag = match.group(3).strip() if match else ""
                message = match.group(4) if match else line
                sensor = SENSORS.search(message)
                if sensor:
                    roll, pitch, ay, water, flags = sensor.groups()
                    row = [datetime.now().astimezone().isoformat(timespec="milliseconds"), level, tag,
                           "IMU_AGUA", number(roll), number(pitch), number(ay), number(water), flags.upper(),
                           None, message, line]
                else:
                    row = [datetime.now().astimezone().isoformat(timespec="milliseconds"), level, tag,
                           "OTRA_LINEA", None, None, None, None, None, message, message, line]
                sheet.append(row)
                print(line)
                if monotonic() - last_save >= SAVE_SECONDS:
                    book.save(output)
                    last_save = monotonic()
    except KeyboardInterrupt:
        print("\nRegistro detenido por el usuario.")
    except serial.SerialException as error:
        print(f"No se pudo abrir el puerto {args.puerto}: {error}")
    finally:
        book.save(output)
        print(f"Archivo Excel guardado: {output}")


if __name__ == "__main__":
    main()
