"""
Lee la salida CSV de una IMU del ESP32 y la guarda en Excel.

Formato esperado del firmware:
    Roll_1,Ay,LTR_Calculado
    12.35,0.42,0.1834

Uso:  py imu_logger_1.py --port COM5 --output imu_prueba_01.xlsx
Notas: 
    - El archivo Excel se crea nuevo; no se sobrescribe uno existente.
    - Es necesario ubicarse correctamente y usar bien las direcciones tanto del .py como del Excel de salida.
"""
from __future__ import annotations
import argparse
import re
from pathlib import Path
import serial
from serial.serialutil import SerialException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


# Columnas del archivo Excel. No se agrega fecha ni hora, como se solicitó.
COLUMNS = ["numero_lectura", "roll_1_deg", "ay_m_s2", "ltr_calculado", "linea_original"]
SAVE_EVERY = 10  # El archivo se actualiza siempre cada 10 lecturas.
# El monitor puede agregar códigos ANSI de color antes de la línea recibida.
ANSI_RE = re.compile(r"\x1B\[[0-?]*[ -/]*[@-~]")

# Convierte los bytes del puerto serial en texto limpio.
def limpiar_linea(raw: bytes) -> str:
    return ANSI_RE.sub("", raw.decode("ascii", errors="replace")).strip()
# Indica si la línea corresponde al encabezado CSV de una IMU.
def es_encabezado(linea: str) -> bool:
    normalizada = linea.replace(" ", "").lower()
    return normalizada == "roll_1,ay,ltr_calculado"

# Convierte una línea CSV numérica en Roll, Ay y LTR.
def leer_muestra(linea: str) -> tuple[float, float, float] | None:
    campos = [campo.strip() for campo in linea.split(",")]
    if len(campos) != 3:
        return None
    try:
        return float(campos[0]), float(campos[1]), float(campos[2])
    except ValueError:
        return None

def crear_libro(path: Path):
    # Crea el libro Excel y prepara la hoja de una IMU.
    if path.exists():
        raise FileExistsError(f"El archivo ya existe: {path}. Use otro nombre para no sobrescribir una prueba.")

    book = Workbook()
    sheet = book.active
    sheet.title = "IMU_1"
    sheet.append(COLUMNS)

    # Formato sencillo para identificar los encabezados y conservar la línea original.
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E1"
    for column, width in { "A": 12, "B": 16,  "C": 14,  "D": 16, "E": 36, }.items():
        sheet.column_dimensions[column].width = width

    return book, sheet


def main() -> None:
    # Lee el puerto COM, filtra el CSV de una IMU y guarda el Excel.
    # argparse solo administra los parámetros necesarios para ejecutar el logger.
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--port", required=True, help="Puerto COM del ESP32, por ejemplo COM5")
    parser.add_argument("--baud", type=int, default=115200, help="Velocidad serial; por defecto 115200")
    parser.add_argument( "--output", default="imu_prueba_01.xlsx",help="Archivo Excel de salida; se crea uno nuevo",)
    args = parser.parse_args()

    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    book, sheet = crear_libro(output)
    book.save(output)

    numero_lectura = 0
    pendientes = 0
    encabezado_encontrado = False
    print(f"Leyendo {args.port} a {args.baud} baudios")
    print(f"Guardando una IMU en {output}")
    print("Presione Ctrl+C para detener y guardar la última parte.")

    try:
        # El ESP32 imprime por UART0/USB; Python recibe ese puerto como COM5, COM6, etc.
        with serial.Serial(args.port, args.baud, timeout=1) as port:
            while True:
                raw = port.readline()
                if not raw:
                    continue

                linea = limpiar_linea(raw)
                if not linea:
                    continue

                if es_encabezado(linea):
                    encabezado_encontrado = True
                    print("Encabezado de una IMU detectado.")
                    continue

                valores = leer_muestra(linea)
                if valores is None:
                    # Ignora ESP_LOG, mensajes de arranque y cualquier línea no CSV.
                    continue

                roll, ay, ltr = valores
                numero_lectura += 1
                sheet.append([numero_lectura, roll, ay, ltr, linea])
                pendientes += 1
                print(linea)

                if pendientes >= SAVE_EVERY:
                    book.save(output)
                    pendientes = 0

    except SerialException as error:
        print(f"No se pudo abrir {args.port}: {error}")
        print("Cierre ESP-IDF Monitor, Excel u otra aplicación que use ese puerto o archivo.")
    except KeyboardInterrupt:
        book.save(output)
        print(f"Registro guardado en {output}")
        if not encabezado_encontrado:
            print("Advertencia: no se detectó el encabezado esperado de una IMU.")


if __name__ == "__main__":
    main()
