"""
Lee el mensaje de ESP-CISTERNA y guarda dos IMUs en Excel.

Formato esperado dentro de la línea del ESP32:
    IMU1[roll=..., pitch=..., ay=...] IMU2[roll=..., pitch=..., ay=...] estado=0x..

Uso:
    py imu_cisterna_logger.py --port COM5 --output "..\\registros\\cisterna_prueba_01.xlsx"
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import serial
from serial.serialutil import SerialException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


# Se conserva una fila por IMU y se identifica con imu_id.
IMU_COLUMNS = [
    "numero_lectura",
    "imu_id",
    "roll_deg",
    "pitch_deg",
    "ay_m_s2",
    "estado",
    "linea_original",
]

SAVE_EVERY = 10

# Busca la parte útil aunque ESP-IDF agregue fecha, etiqueta, colores o nivel.
MESSAGE_RE = re.compile(
    r"IMU1\[roll=(?P<roll1>[-+0-9.eE]+),\s*"
    r"pitch=(?P<pitch1>[-+0-9.eE]+),\s*"
    r"ay=(?P<ay1>[-+0-9.eE]+)\]\s*"
    r"IMU2\[roll=(?P<roll2>[-+0-9.eE]+),\s*"
    r"pitch=(?P<pitch2>[-+0-9.eE]+),\s*"
    r"ay=(?P<ay2>[-+0-9.eE]+)\]\s*"
    r"(?:nivel=[-+0-9.eE]+\s*)?"
    r"estado=0x(?P<status>[0-9A-Fa-f]+)"
)

ANSI_RE = re.compile(r"\x1B\[[0-?]*[ -/]*[@-~]")

#Convierte los bytes del puerto serial en texto sin códigos ANSI.
def limpiar_linea(raw: bytes) -> str:
    return ANSI_RE.sub("", raw.decode("ascii", errors="replace")).strip()

# Extrae ambas IMUs y el estado de una línea del ESP32.
def leer_mensaje(linea: str) -> dict[str, float | str] | None:
    match = MESSAGE_RE.search(linea)
    if match is None:
        return None

    valores = match.groupdict()
    return {
        "roll1": float(valores["roll1"]),
        "pitch1": float(valores["pitch1"]),
        "ay1": float(valores["ay1"]),
        "roll2": float(valores["roll2"]),
        "pitch2": float(valores["pitch2"]),
        "ay2": float(valores["ay2"]),
        "status": f"0x{valores['status'].upper()}",
    }

# Aplica encabezados y formato básico a una hoja Excel.
def aplicar_encabezado(sheet, columns: list[str], widths: dict[str, int]) -> None:
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{chr(64 + len(columns))}1"
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

# Crea el Excel con una tabla normalizada para ambas IMUs.
def crear_libro(path: Path):
    if path.exists():
        raise FileExistsError(f"El archivo ya existe: {path}. Use otro nombre.")

    book = Workbook()
    imus = book.active
    imus.title = "IMUs"

    aplicar_encabezado(
        imus,
        IMU_COLUMNS,
        {"A": 16, "B": 12, "C": 14, "D": 14, "E": 14, "F": 12, "G": 72},
    )
    return book, imus

# Lee el COM y registra cada mensaje completo de ESP-CISTERNA.
def main() -> None:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--port", required=True)
    parser.add_argument("--baud", type=int, default=115200)
    parser.add_argument("--output", default="cisterna_prueba_01.xlsx")
    args = parser.parse_args()

    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    book, imus = crear_libro(output)
    book.save(output)

    numero_lectura = 0
    pendientes = 0
    print(f"Leyendo ESP-CISTERNA en {args.port} a {args.baud} baudios")
    print(f"Guardando en {output}")
    print("Presione Ctrl+C para detener y guardar la última parte.")

    try:
        with serial.Serial(args.port, args.baud, timeout=1) as port:
            while True:
                raw = port.readline()
                if not raw:
                    continue

                linea = limpiar_linea(raw)
                datos = leer_mensaje(linea)
                if datos is None:
                    continue

                numero_lectura += 1
                imus.append(
                    [
                        numero_lectura,
                        "IMU1",
                        datos["roll1"],
                        datos["pitch1"],
                        datos["ay1"],
                        datos["status"],
                        linea,
                    ]
                )
                imus.append(
                    [
                        numero_lectura,
                        "IMU2",
                        datos["roll2"],
                        datos["pitch2"],
                        datos["ay2"],
                        datos["status"],
                        linea,
                    ]
                )
                pendientes += 1
                print(linea)

                if pendientes >= SAVE_EVERY:
                    book.save(output)
                    pendientes = 0

    except SerialException as error:
        print(f"No se pudo abrir {args.port}: {error}")
        print("Cierre ESP-IDF Monitor, Excel u otra aplicación que use ese puerto.")
    except KeyboardInterrupt:
        book.save(output)
        print(f"Registro guardado en {output}")


if __name__ == "__main__":
    main()
